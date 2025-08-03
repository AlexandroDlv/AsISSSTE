from django.shortcuts import render, redirect
import sys, os
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import db_con
from datetime import datetime, timedelta
import calendar
from django.contrib.auth import authenticate, login
from django.contrib import messages
from openpyxl import Workbook # type: ignore
from openpyxl.styles import Font # type: ignore
from openpyxl.utils import get_column_letter # type: ignore son para quitar los avisos de error
from django.http import HttpResponse

# Create your views here.

def inicio(request):
    return render(request, 'inicio.html')

def principal(request):
    if 'usuario' not in request.session:
        return redirect('login')

    nombre = request.session.get('nombre', 'Usuario')
    return render(request, 'principal.html', {'nombre': nombre})

def inicia_sesion(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')

        personal_col = db_con.db["personal"]
        usuario = personal_col.find_one({"usuario": username})

        if usuario and usuario["contraseña"] == password:
            request.session['usuario'] = usuario["usuario"]
            request.session['nombre'] = usuario["nombre"]
            return redirect('principal')
        else:
            messages.error(request, 'Usuario o contraseña incorrectos')

    return render(request, 'login.html')

def convertir_actividades(actividades):
    lista = []
    for act in actividades:
        lista.append({
            "nombre": act.get("nombre", ""),
            "horarios": act.get("horarios", [])
        })
    return lista

def registro_derechohabiente(request):
    actividades_col = db_con.db["actividades"]
    personas_col = db_con.db["personas"]

    if request.method == "POST":
        actividades_seleccionadas = request.POST.getlist("actividad")
        horarios_seleccionados = request.POST.getlist("horario")

        actividades_y_horarios = []
        for act, hor in zip(actividades_seleccionadas, horarios_seleccionados):
            actividades_y_horarios.append({
                "actividad": act,
                "horario": hor
            })

        datos = {
            "id_Personal": generar_id_personal(),
            "nombre": request.POST.get("nombre"),
            "apellido": request.POST.get("apellido"),
            "edad": int(request.POST.get("edad")),
            "curp": request.POST.get("curp"),
            "telefono": request.POST.get("telefono"),
            "telefono_emergencia": request.POST.get("telefono_emergencia"),
            "correo": request.POST.get("correo"),
            "fecha_registro": datetime.now().strftime("%Y-%m-%d"),
            "genero": request.POST.get("genero"),
            "actividades": actividades_y_horarios 
        }
        personas_col.insert_one(datos)
        return redirect("vista_personas")

    actividades = list(actividades_col.find())
    actividades_json = convertir_actividades(actividades)
    return render(request, "registro.html", {"actividades": actividades_json})

def vista_personas(request):
    personas_col = db_con.db["personas"]
    query = request.GET.get("buscar", "").strip()
    
    if query:
        datos = list(personas_col.find(
            {"nombre": {"$regex": query, "$options": "i"}},
            {'_id': 0}
        ))
    else:
        datos = list(personas_col.find({}, {'_id': 0}))

    return render(request, 'personas.html', {
        'personas': datos,
        'buscar': query
    })

def vista_actividades(request):
    actividades_col = db_con.db["actividades"]
    actividades = list(actividades_col.find({}, {'_id': 0}))
    return render(request, 'actividades.html', {'actividades': actividades})

def generar_id_personal():
    personas_col = db_con.db["personas"]
    ultima = personas_col.find_one(sort=[("id_Personal", -1)])
    if ultima and "id_Personal" in ultima:
        try:
            ultimo_numero = int(ultima["id_Personal"].split("-")[-1])
        except:
            ultimo_numero = 0
    else:
        ultimo_numero = 0
    nuevo_id = f"PER-{ultimo_numero + 1:03}"
    return nuevo_id

def exportar_asistencia_excel(request):
    personas_col = db_con.db["personas"]
    hoy = datetime.now()
    mes_actual = hoy.strftime("%B %Y")
    ultimo_dia = calendar.monthrange(hoy.year, hoy.month)[1]
    dias_mes = [f"{dia:02}" for dia in range(1, ultimo_dia + 1)]

    personas = list(personas_col.find({}, {"_id": 0}))
    asistencia_por_actividad_horario = {}

    for persona in personas:
        actividades = persona.get("actividades", [])
        for act in actividades:
            nombre_actividad = act.get("actividad", "Sin Actividad")
            horario = act.get("horario", "Sin Horario")
            clave = f"{nombre_actividad} - {horario}"

            if clave not in asistencia_por_actividad_horario:
                asistencia_por_actividad_horario[clave] = []

            asistencia_por_actividad_horario[clave].append(persona)

    #archivo excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Asistencia"

    fila_actual = 1
    for clave, personas in asistencia_por_actividad_horario.items():
        ws.merge_cells(start_row=fila_actual, start_column=1, end_row=fila_actual, end_column=3 + len(dias_mes))
        celda_titulo = ws.cell(row=fila_actual, column=1)
        celda_titulo.value = f"Asistencia - {clave} - {mes_actual}"
        celda_titulo.font = Font(bold=True, size=14)
        fila_actual += 1

        encabezados = ["ID Personal", "Nombre", "Apellidos"] + dias_mes
        for col, encabezado in enumerate(encabezados, start=1):
            celda = ws.cell(row=fila_actual, column=col)
            celda.value = encabezado
            celda.font = Font(bold=True)
        fila_actual += 1

        for persona in personas:
            ws.cell(row=fila_actual, column=1).value = persona.get("id_Personal", "")
            ws.cell(row=fila_actual, column=2).value = persona.get("nombre", "")
            ws.cell(row=fila_actual, column=3).value = persona.get("apellido", "")
            for idx in range(4, 4 + len(dias_mes)):
                ws.cell(row=fila_actual, column=idx).value = ""
            fila_actual += 1

        fila_actual += 2 

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    nombre_archivo = f"asistencia_{hoy.strftime('%Y_%m')}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{nombre_archivo}"'
    anchos = [15, 20, 25] + [5] * len(dias_mes)  # Ajusta según el contenido que esperas

    for i, ancho in enumerate(anchos, start=1):
        col_letra = get_column_letter(i)
        ws.column_dimensions[col_letra].width = ancho
    wb.save(response)
    return response

def lista_asistencia(request):
    personas_col = db_con.db["personas"]
    asistencia_col = db_con.db["asistencia"]

    hoy = datetime.now()
    mes_actual = hoy.strftime("%B %Y")
    ultimo_dia = calendar.monthrange(hoy.year, hoy.month)[1]
    dias_mes = [f"{dia:02}" for dia in range(1, ultimo_dia + 1)]

    personas = list(personas_col.find({}, {"_id": 0}))
    asistencia_por_actividad = {}

    for persona in personas:
        actividad = persona.get("actividad", "Sin Actividad")
        if actividad not in asistencia_por_actividad:
            asistencia_por_actividad[actividad] = []

        asistencias = asistencia_col.find({
            "id_Personal": persona["id_Personal"],
            "fecha": {"$regex": f"^{hoy.strftime('%Y-%m')}"}
        })
        asistencias_dict = {a["fecha"][-2:]: a["asistio"] for a in asistencias}
        persona["asistencias"] = asistencias_dict

        asistencia_por_actividad[actividad].append(persona)

    return render(request, "asistencia.html", {
        "asistencia_por_actividad": asistencia_por_actividad,
        "dias_mes": dias_mes,
        "mes_actual": mes_actual
    })

def generar_id_act():
    actividades_col = db_con.db["actividades"]
    ultima = actividades_col.find().sort("id_actividad", -1).limit(1)
    try:
        ultimo = list(ultima)[0]['id_actividad']
        numero = int(ultimo.split("-")[1]) + 1
    except:
        numero = 1
    return f"ACT-{numero:03d}"

def registrar_actividad(request):
    actividades_col = db_con.db["actividades"]
    if request.method == "POST":
        nombre = request.POST.get("nombre")
        horarios = request.POST.getlist("horarios")

        nueva_actividad = {
            "id_actividad": generar_id_act(),
            "nombre": nombre,
            "horarios": horarios
        }

        actividades_col.insert_one(nueva_actividad)
        return redirect('vista_actividades') 

    return render(request, "reg_actividad.html")

def editar_actividad(request, id_actividad):
    actividades_col = db_con.db["actividades"]
    actividad = actividades_col.find_one({"id_actividad": id_actividad})

    if not actividad:
        return redirect("vista_actividades")

    if request.method == "POST":
        nuevo_nombre = request.POST.get("nombre")
        nuevos_horarios = request.POST.getlist("horarios")
        nuevos_horarios = [h.strip() for h in nuevos_horarios if h.strip()] 

        actividades_col.update_one(
            {"id_actividad": id_actividad},
            {"$set": {
                "nombre": nuevo_nombre,
                "horarios": nuevos_horarios
            }}
        )
        return redirect("vista_actividades")

    return render(request, "editar_actividad.html", {"actividad": actividad})

def eliminar_actividad(request, id_actividad):
    actividades_col = db_con.db["actividades"]
    actividades_col.delete_one({'id_actividad': id_actividad})
    return redirect('vista_actividades')

def editar_persona(request, id_personal):
    personas_col = db_con.db["personas"]
    actividades_col = db_con.db["actividades"]
    persona = personas_col.find_one({"id_Personal": id_personal})

    if not persona:
        return redirect("vista_personas")

    if request.method == "POST":
        actividades_seleccionadas = request.POST.getlist("actividad")
        horarios_seleccionados = request.POST.getlist("horario")

        actividades_y_horarios = []
        for act, hor in zip(actividades_seleccionadas, horarios_seleccionados):
            actividades_y_horarios.append({
                "actividad": act,
                "horario": hor
            })

        datos_actualizados = {
            "nombre": request.POST.get("nombre"),
            "apellido": request.POST.get("apellido"),
            "edad": int(request.POST.get("edad")),
            "curp": request.POST.get("curp"),
            "telefono": request.POST.get("telefono"),
            "telefono_emergencia": request.POST.get("telefono_emergencia"),
            "correo": request.POST.get("correo"),
            "genero": request.POST.get("genero"),
            "actividades": actividades_y_horarios
        }

        personas_col.update_one(
            {"id_Personal": id_personal},
            {"$set": datos_actualizados}
        )
        return redirect("vista_personas")

    actividades = list(actividades_col.find({}, {"_id": 0}))
    return render(request, "editar_persona.html", {
        "persona": persona,
        "actividades": actividades
    })

def eliminar_persona(request, id_personal):
    personas_col = db_con.db["personas"]
    persona = personas_col.find_one({"id_Personal": id_personal})

    if persona:
        personas_col.delete_one({"id_Personal": id_personal})

    return redirect("vista_personas")